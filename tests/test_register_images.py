from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from scripts.common import sha
from scripts.register_images import delete_column, register, rows
from tests.test_pipeline import run


def photo(ctx, name, color="red"):
    Image.new("RGB", (23, 41), color).save(ctx.originals / name)


def test_deleted_column_preserves_retained_validation():
    ws = Workbook().active
    for target in ["D2:D5", "E2:E5", "F2:F5"]:
        validation = DataValidation(type="list", formula1='"a,b"')
        validation.add(target)
        ws.add_data_validation(validation)
    delete_column(ws, 5)
    delete_column(ws, 4)
    assert [str(v.sqref) for v in ws.data_validations.dataValidation] == ["D2:D5"]


def test_preview_apply_and_repeat_preserve_metadata(ctx):
    photo(ctx, "TX000001_MO0002_close_up.jpg")
    before = sha(ctx.master)
    assert run(ctx, register).ok
    assert sha(ctx.master) == before
    assert run(ctx, register, apply=True).ok
    wb = load_workbook(ctx.master)
    images = rows(wb["06_Images"])
    assert images[-1]["image_id"] == "IM000003"
    assert images[-1]["original_filename"] == "TX000001_MO0002_close_up.jpg"
    assert images[0]["note"] == "DO_NOT_PUBLISH"
    assert ("IM000003", "MO0002") in {(r["image_id"], r["morphology_id"]) for r in rows(wb["07_Image_Morphology"])}
    assert len(rows(wb["05_Taxon_Morphology"])) == 2
    wb.close()
    saved = sha(ctx.master)
    assert run(ctx, register, apply=True).ok
    assert sha(ctx.master) == saved
    assert len(list((ctx.root / "backups").glob("*.xlsx"))) == 1


def test_invalid_batch_is_atomic(ctx):
    photo(ctx, "TX000001_MO0001_good.jpg")
    photo(ctx, "TX999999_MO0001_invalid.jpg", "blue")
    before = sha(ctx.master)
    assert not run(ctx, register, apply=True).ok
    assert sha(ctx.master) == before


def test_duplicate_photo_requires_reconciliation(ctx):
    (ctx.originals / "TX000001_MO0001_copy.jpg").write_bytes((ctx.originals / "first.jpg").read_bytes())
    before = sha(ctx.master)
    assert not run(ctx, register, apply=True).ok
    assert sha(ctx.master) == before


def test_migration_preserves_remaining_cells(ctx):
    wb = load_workbook(ctx.master)
    ws = wb["06_Images"]
    ws.cell(1, ws.max_column + 1, "status")
    ws.cell(2, ws.max_column, "active")
    before = rows(ws)[0]
    wb.save(ctx.master)
    wb.close()
    assert not run(ctx, register).ok
    assert run(ctx, register, migrate=True, apply=True).ok
    wb = load_workbook(ctx.master)
    after = rows(wb["06_Images"])[0]
    assert after == {k: v for k, v in before.items() if k != "status"}
    wb.close()


def test_renamed_registered_file_is_not_assigned_new_id(ctx):
    (ctx.originals / "first.jpg").rename(ctx.originals / "TX000001_MO0001_renamed.jpg")
    assert not run(ctx, register, apply=True).ok
