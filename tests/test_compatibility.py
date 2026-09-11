from openpyxl import load_workbook

from scripts.common import Report
from scripts.upload_images import preflight
from scripts.validate import validate_source


def test_only_explicit_unnamed_column_can_be_ignored(ctx):
    wb = load_workbook(ctx.master)
    sheet = wb["01_Taxa"]
    col = sheet.max_column + 1
    sheet.cell(2, col, "uninterpreted existing value")
    letter = sheet.cell(1, col).column_letter
    wb.save(ctx.master)
    ctx.config["ignored_unnamed_columns"] = {}
    report = Report(ctx, "source")
    validate_source(ctx, report)
    assert not report.ok
    ctx.config["ignored_unnamed_columns"] = {"01_Taxa": [letter]}
    report = Report(ctx, "source")
    validate_source(ctx, report)
    assert report.ok
    assert "UNNAMED_COLUMN_PRESERVED" in {issue["code"] for issue in report.issues}


def test_test_dataset_is_blocked_before_remote_access(ctx):
    ctx.config["test_dataset"] = {"enabled": True}
    report = Report(ctx, "upload")
    assert preflight(ctx, report) == (None, {})
    assert not report.ok
    assert report.issues[0]["code"] == "TEST_DATASET_LOCAL_ONLY"
