import shutil

import pytest
from openpyxl import load_workbook
from PIL import Image

from scripts.common import ROOT, Context, Report
from scripts.init_master import initialize


def append(ctx, sheet, **row):
    wb = load_workbook(ctx.master)
    ws = wb[sheet]
    ws.append([row.get(c.value) for c in ws[1]])
    wb.save(ctx.master)
    wb.close()


@pytest.fixture
def ctx(tmp_path):
    shutil.copytree(ROOT / "config", tmp_path / "config")
    c = Context(tmp_path)
    c.config["test_dataset"] = {"enabled": False}
    initialize(c, Report(c, "init"))
    append(
        c,
        "01_Taxa",
        taxon_id="TX000001",
        scientific_name="Synthetic taxon A",
        taxon_rank="SPECIES",
        korean_name="합성 분류군",
    )
    append(c, "01_Taxa", taxon_id="TX000002", scientific_name="Synthetic taxon B", taxon_rank="SPECIES")
    append(c, "03_Morphology_Terms", term_id="MT0001", english="Synthetic term", category="general")
    append(c, "04_Morphology_Tree", morphology_id="MO0001", term_id="MT0001", display_order=0)
    append(
        c, "04_Morphology_Tree", morphology_id="MO0002", parent_id="MO0001", term_id="MT0001", display_order=1
    )
    for ident, filename, status in [
        ("IM000001", "first.jpg", "active"),
        ("IM000002", "sub/second.jpg", "hidden"),
    ]:
        path = c.originals / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (800, 400), (40, 110, 70)).save(path)
        append(
            c,
            "06_Images",
            image_id=ident,
            taxon_id="TX000001",
            original_filename=filename,
            photographer="Synthetic author",
            copyright="Synthetic test only",
            status=status,
            note="DO_NOT_PUBLISH",
            locality="SECRET_PRECISE_LOCALITY",
        )
    append(
        c,
        "07_Image_Morphology",
        image_id="IM000001",
        morphology_id="MO0001",
        role="primary",
        representative=True,
        display_order=0,
    )
    append(
        c,
        "08_Synonyms",
        name_id="NA0001",
        taxon_id="TX000001",
        name="Synthetic synonym",
        name_type="scientific_synonym",
    )
    return c


def change(ctx, sheet, column, value, row=2):
    wb = load_workbook(ctx.master)
    ws = wb[sheet]
    col = [c.value for c in ws[1]].index(column) + 1
    ws.cell(row, col, value)
    wb.save(ctx.master)
    wb.close()
