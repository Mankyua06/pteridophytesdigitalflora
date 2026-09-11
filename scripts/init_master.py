from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .common import cli, safe_path


def initialize(ctx, report, filename=None):
    path = safe_path(ctx.root, filename) if filename else ctx.master
    if path.exists():
        report.add("error", "MASTER_EXISTS", "기존 Excel을 덮어쓰지 않습니다.", path=str(path))
        return
    wb = Workbook()
    wb.remove(wb.active)
    for name, spec in ctx.schema.items():
        ws = wb.create_sheet(name)
        ws.append(list(spec["columns"]))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell, (column, rule) in zip(ws[1], spec["columns"].items()):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244B39")
            ws.column_dimensions[cell.column_letter].width = max(20, len(column) + 3)
            if rule["type"] in ("id", "external_key", "string"):
                ws.column_dimensions[cell.column_letter].number_format = "@"
            if rule.get("enum") or rule["type"] == "boolean":
                values = rule.get("enum", ["TRUE", "FALSE"])
                dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
                dv.errorTitle, dv.error, dv.showErrorMessage = "허용값 확인", "목록의 값을 입력하세요.", True
                ws.add_data_validation(dv)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}1048576")
    # Exclusive creation protects against a file appearing after the exists check.
    with path.open("xb") as f:
        wb.save(f)
    for name in (
        "images_original",
        "images_web/thumb",
        "images_web/medium",
        "images_web/large",
        ".state",
        "reports",
        "backups",
        "work",
    ):
        (ctx.root / name).mkdir(parents=True, exist_ok=True)
    report.counts["generated"] += 1


if __name__ == "__main__":
    cli(
        "init",
        lambda p: p.add_argument("--filename", help="새 빈 템플릿의 상대경로"),
        lambda c, r, a: initialize(c, r, a.filename),
    )
