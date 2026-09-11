import re

from openpyxl import load_workbook

from .common import read_json


def resolve(ctx, rows, language='en'):
    defaults = read_json(ctx.root / 'config/site_text.json', {})
    translations = read_json(ctx.root / 'config/site_text_ko.json', {})
    result = dict(defaults if language == 'en' else translations)
    seen = set()
    for row in rows:
        key, text = row['key'], row['text']
        if key not in defaults or key in seen or not isinstance(text, str) or not text.strip():
            raise ValueError(f'Invalid or duplicate site text key/value: {key}')
        if set(re.findall(r'\{\w+\}', text)) != set(re.findall(r'\{\w+\}', defaults[key])):
            raise ValueError(f'Placeholders must be preserved: {key}')
        seen.add(key)
        korean = row.get('text_ko')
        if korean is not None and (not isinstance(korean, str) or (korean.strip() and
                set(re.findall(r'\{\w+\}', korean)) != set(re.findall(r'\{\w+\}', defaults[key])))):
            raise ValueError(f'Korean placeholders must be preserved: {key}')
        if language == 'en':
            result[key] = text
        else:
            # An explicitly blank translation uses the current English override.
            result[key] = korean if isinstance(korean, str) and korean.strip() else text
    return result


def read_master(ctx):
    wb = load_workbook(ctx.master, read_only=True, data_only=False)
    try:
        if '10_Site_Text' not in wb.sheetnames:
            return {'text': resolve(ctx, []), 'text_ko': resolve(ctx, [], 'ko')}
        sheet = wb['10_Site_Text']
        header = [c.value for c in sheet[1]]
        if header.count('key') != 1 or header.count('text') != 1:
            raise ValueError('Site text needs unique key and text columns')
        rows = []
        for cells in sheet.iter_rows(min_row=2):
            if all(c.value is None for c in cells):
                continue
            values = {k:c for k,c in zip(header,cells)}
            if any(values[k].data_type == 'f' for k in ('key','text','text_ko') if k in values):
                raise ValueError('Site text formulas are not supported')
            rows.append({k:values[k].value for k in ('key','text','text_ko') if k in values})
        return {'text': resolve(ctx, rows), 'text_ko': resolve(ctx, rows, 'ko')}
    finally:
        wb.close()
