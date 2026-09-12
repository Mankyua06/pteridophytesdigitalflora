from __future__ import annotations

import math
import re
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook
from PIL import Image

from .common import ID_RE, Context, Report, digest, parser, read_json, safe_path, sha


def normalize(value, rule):
    if value is None or value == "":
        return None
    kind = rule["type"]
    if kind == "id":
        if not isinstance(value, str) or not ID_RE.fullmatch(value):
            raise ValueError("ID는 안전한 문자열이어야 합니다. 숫자 ID를 자동 변환하지 않습니다.")
    elif kind == "string":
        if not isinstance(value, str):
            raise ValueError("텍스트 형식이 필요합니다. 원문을 텍스트로 명시하세요.")
    elif kind == "external_key":
        if isinstance(value, bool) or not isinstance(value, (str, int)):
            raise ValueError("외부 키는 문자열 또는 정수여야 합니다.")
        value = str(value)
    elif kind == "boolean":
        if isinstance(value, bool):
            pass
        elif isinstance(value, str) and value in ("TRUE", "FALSE"):
            value = value == "TRUE"
        else:
            raise ValueError("실제 boolean 또는 TRUE/FALSE만 허용합니다.")
    elif kind in ("integer", "positive_number"):
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError("유한 숫자가 필요합니다.")
        if kind == "integer":
            if value != int(value) or value < rule.get("minimum", 0):
                raise ValueError("0 이상의 정수가 필요합니다.")
            value = int(value)
        elif value <= 0:
            raise ValueError("양수여야 합니다.")
    elif kind == "date":
        if isinstance(value, datetime):
            value = value.date().isoformat()
        elif isinstance(value, date):
            value = value.isoformat()
        elif isinstance(value, str):
            parsed = datetime.strptime(value, "%Y-%m-%d").date()
            if parsed.isoformat() != value:
                raise ValueError("날짜 문자열은 YYYY-MM-DD만 허용합니다.")
            value = parsed.isoformat()
        else:
            raise ValueError("Excel 날짜 셀 또는 YYYY-MM-DD만 허용합니다.")
    if rule.get("enum") and value not in rule["enum"]:
        raise ValueError(f"허용값: {', '.join(rule['enum'])}")
    return value


def load_source(ctx, report):
    data = {s: [] for s in ctx.schema}
    if not ctx.master.exists():
        report.add("error", "MASTER_MISSING", "설정된 master Excel이 없습니다.", path=str(ctx.master))
        return data
    wb = load_workbook(ctx.master, data_only=False)
    for name, spec in ctx.schema.items():
        report.counts["checks"] += 1
        if name not in wb.sheetnames:
            optional = name in ctx.config.get("optional_sheets", [])
            report.add(
                "warning" if optional else "error",
                "OPTIONAL_SHEET_MISSING" if optional else "SHEET_MISSING",
                "설정상 선택 시트가 없어 빈 목록으로 읽습니다." if optional else "필수 시트가 없습니다.",
                sheet=name,
            )
            continue
        ws = wb[name]
        if ws.merged_cells.ranges:
            report.add("error", "MERGED_CELLS", "병합 셀을 해제하세요.", sheet=name)
        aliases = ctx.config.get("aliases", {}).get(name, {})
        headers = [aliases.get(c.value, c.value) for c in ws[1]]
        ignored_unnamed = set(ctx.config.get("ignored_unnamed_columns", {}).get(name, []))
        for cell in ws[1]:
            if cell.value is None and cell.column_letter in ignored_unnamed:
                report.add(
                    "warning",
                    "UNNAMED_COLUMN_PRESERVED",
                    "명시적 호환 설정: 원본 값은 보존하고 의미 해석·공개를 제외합니다.",
                    sheet=name,
                    column=cell.column_letter,
                )
        populated = [h for h in headers if h is not None]
        if len(populated) != len(set(populated)):
            report.add("error", "DUPLICATE_HEADER", "중복 열 이름이 있습니다.", sheet=name)
        for h in spec["columns"]:
            if h not in headers:
                report.add(
                    "error" if h in spec["required"] else "warning",
                    "COLUMN_MISSING",
                    "필수 열 누락" if h in spec["required"] else "선택 열 누락: 내부 null 처리",
                    sheet=name,
                    column=h,
                )
        for h in populated:
            if h not in spec["columns"]:
                report.add(
                    "warning",
                    "EXTRA_COLUMN",
                    "추가 열 보존. 공개 허용 목록에 없으면 제외합니다.",
                    sheet=name,
                    column=h,
                )
            if name == "06_Images" and (
                h == "filename"
                or "url" in str(h).lower()
                or h in ("thumb_filename", "medium_filename", "large_filename")
            ):
                report.add(
                    "warning",
                    "DERIVED_COLUMN_IGNORED",
                    "기존 파생 열을 경로 결정에 사용하지 않습니다.",
                    sheet=name,
                    column=h,
                )
        for cells in ws.iter_rows(min_row=2):
            if all(c.value is None for c in cells):
                continue
            row = {h: None for h in spec["columns"]}
            row["_row"] = cells[0].row
            for h, cell in zip(headers, cells):
                report.counts["checks"] += 1
                loc = dict(sheet=name, row=cell.row, column=h or cell.column_letter)
                if cell.data_type == "f":
                    # Extra helper formulas have no authority and are never evaluated/exported.
                    report.add(
                        "error" if h in spec["columns"] else "warning",
                        "FORMULA",
                        "수식을 평가하지 않습니다. 핵심 열은 명시적 값이 필요합니다.",
                        **loc,
                    )
                    continue
                if h is None:
                    if cell.value is not None and cell.column_letter not in ignored_unnamed:
                        report.add(
                            "error",
                            "UNNAMED_COLUMN",
                            "제목 없는 열에 데이터가 있습니다. 의미를 확인해 이름을 지정하세요.",
                            **loc,
                        )
                    continue
                try:
                    row[h] = normalize(cell.value, spec["columns"].get(h, {"type": "string"}))
                except (ValueError, TypeError):
                    report.add(
                        "error",
                        "INVALID_VALUE",
                        f"유효하지 않은 {spec['columns'].get(h, {}).get('type', 'string')} 값",
                        **loc,
                    )
            for h in spec["required"]:
                if row.get(h) is None:
                    report.add(
                        "error", "REQUIRED", "필수값이 없습니다.", sheet=name, row=row["_row"], column=h
                    )
            data[name].append(row)
    wb.close()
    return data


def validate_source(ctx, report):
    data = load_source(ctx, report)
    for sheet, columns in [("01_Taxa", ("scientific_name", "accepted_scientific_name")),
                           ("08_Synonyms", ("name",))]:
        for row in data.get(sheet, []):
            for column in columns:
                opened = False
                valid = True
                for marker in re.findall(r"\[/?i\]", row.get(column) or ""):
                    if (marker == "[i]") == opened:
                        valid = False
                    opened = marker == "[i]"
                if opened or not valid:
                    report.add("error", "SCIENTIFIC_NAME_MARKUP", "[i]와 [/i]를 중첩 없이 짝지어 입력하세요.",
                               sheet=sheet, row=row["_row"], column=column)
    from .site_text import resolve
    try:
        resolve(ctx, data.get("10_Site_Text", []))
    except ValueError as error:
        report.add("error", "SITE_TEXT_INVALID", str(error), sheet="10_Site_Text")
    for name, spec in ctx.schema.items():
        seen = set()
        for row in data[name]:
            key = tuple(row.get(k) for k in spec["primary_key"])
            report.counts["checks"] += 1
            if key in seen:
                report.add(
                    "error", "DUPLICATE_KEY", "중복 기본키", sheet=name, row=row["_row"], record_id=str(key)
                )
            seen.add(key)
            for column, rule in spec["columns"].items():
                if rule.get("references") and row.get(column) is not None:
                    target, field = rule["references"].split(".")
                    if row[column] not in {r[field] for r in data[target]}:
                        report.add(
                            "error",
                            "FOREIGN_KEY",
                            "외래키 대상이 없습니다.",
                            sheet=name,
                            row=row["_row"],
                            column=column,
                            record_id=row[column],
                        )
    tree = {r["morphology_id"]: r["parent_id"] for r in data["04_Morphology_Tree"] if r["morphology_id"]}
    for node in tree:
        current, visited = node, set()
        while current in tree:
            if current in visited:
                report.add("error", "TREE_CYCLE", "형태 트리 순환 또는 자기 참조", record_id=node)
                break
            visited.add(current)
            current = tree[current]
    cytotypes = {r["cytotype_id"]: r for r in data["02_Cytotypes"]}
    registered, folded = set(), {}
    for row in data["06_Images"]:
        loc = dict(sheet="06_Images", row=row["_row"], record_id=row["image_id"])
        cy = cytotypes.get(row["cytotype_id"])
        if cy and cy["taxon_id"] != row["taxon_id"]:
            report.add("error", "CYTOTYPE_TAXON", "이미지와 cytotype의 taxon이 다릅니다.", **loc)
        try:
            original = safe_path(ctx.originals, row["original_filename"])
            registered.add(original)
            with Image.open(original) as im:
                im.load()
            report.counts["checks"] += 1
        except (ValueError, TypeError):
            report.add("error", "UNSAFE_PATH", "원본 상대경로가 잘못되었습니다.", **loc)
        except FileNotFoundError:
            report.add(
                "error", "ORIGINAL_MISSING", "등록된 원본이 없습니다.", path=row["original_filename"], **loc
            )
        except Exception:
            report.add(
                "error",
                "ORIGINAL_DECODE",
                "원본 디코딩 실패 또는 지원하지 않는 형식",
                path=row["original_filename"],
                **loc,
            )
        if row["status"] == "active":
            for key in ("photographer", "copyright"):
                if not row.get(key):
                    report.add("warning", "CREDIT_MISSING", "공개 이미지 표시 정보 누락", column=key, **loc)
    for path in ctx.originals.rglob("*"):
        if path.is_file():
            rel = path.relative_to(ctx.originals).as_posix()
            if rel.casefold() in folded and folded[rel.casefold()] != rel:
                report.add("error", "CASE_COLLISION", "원본 경로 대소문자 충돌", path=rel)
            folded[rel.casefold()] = rel
            if path.resolve() not in registered:
                report.add(
                    "warning", "UNREGISTERED_ORIGINAL", "Excel 미등록 원본. 생성 대상이 아닙니다.", path=rel
                )
    for sheet, idcol in (("01_Taxa", "taxon_id"), ("06_Images", "image_id")):
        ids = [r[idcol] for r in data[sheet] if r[idcol]]
        if len({i.casefold() for i in ids}) != len(set(ids)):
            report.add("error", "CASE_COLLISION", "ID 대소문자 충돌", sheet=sheet)
    links = data["07_Image_Morphology"]
    images = {r["image_id"]: r for r in data["06_Images"]}
    primaries = Counter(r["image_id"] for r in links if r["role"] == "primary")
    reps = Counter(
        (images.get(r["image_id"], {}).get("taxon_id"), r["morphology_id"])
        for r in links
        if r["representative"]
    )
    if any(n > 1 for n in primaries.values()):
        report.add("warning", "MULTIPLE_PRIMARY", "한 사진에 primary 연결이 여러 개입니다.")
    if any(n > 1 for n in reps.values()):
        report.add("warning", "MULTIPLE_REPRESENTATIVE", "같은 taxon/형태의 대표사진이 여러 개입니다.")
    for row in data["05_Taxon_Morphology"]:
        matching = [
            images[r["image_id"]]
            for r in links
            if r["image_id"] in images
            and r["morphology_id"] == row["morphology_id"]
            and images[r["image_id"]]["taxon_id"] == row["taxon_id"]
        ]
        if row["documentation_status"] == "documented":
            if not matching:
                report.add(
                    "warning",
                    "DOCUMENTED_NO_IMAGE",
                    "documented이지만 연결 사진이 없습니다.",
                    row=row["_row"],
                )
            elif not any(i["status"] == "active" for i in matching):
                report.add(
                    "warning",
                    "DOCUMENTED_NONPUBLIC_ONLY",
                    "자료는 있으나 공개 사진이 없습니다.",
                    row=row["_row"],
                )
        if matching and row["presence_status"] == "absent":
            report.add(
                "warning",
                "ABSENT_WITH_IMAGE",
                "absent 상태에 연결 사진이 있습니다. 상태는 자동 변경하지 않습니다.",
                row=row["_row"],
            )
    return data


def inspect_webp(path):
    with Image.open(path) as im:
        im.load()
        if im.format != "WEBP":
            raise ValueError("not WebP")
        return dict(sha256=sha(path), width=im.width, height=im.height, bytes=path.stat().st_size)


def validate_derived(ctx, report, data=None, active_only=False):
    if data is None:
        data = validate_source(ctx, report)
    manifest = read_json(ctx.state, {})
    settings = digest(ctx.settings())
    expected = set()
    for row in data["06_Images"]:
        if active_only and row["status"] != "active":
            continue
        ident = row["image_id"]
        if not ident or not ID_RE.fullmatch(ident):
            continue
        entry = manifest.get(ident)
        report.counts["checks"] += 1
        if not entry:
            report.add("error", "MANIFEST_MISSING", "생성 이력 없음. rebuild가 필요합니다.", record_id=ident)
            continue
        try:
            if (
                entry["source"] != row["original_filename"]
                or entry["source_sha256"] != sha(safe_path(ctx.originals, row["original_filename"]))
                or entry["settings"] != settings
            ):
                report.add(
                    "error",
                    "STALE",
                    "원본 또는 설정이 변경되었습니다. rebuild가 필요합니다.",
                    record_id=ident,
                )
            for size, limit in ctx.config["sizes"].items():
                rel = f"{size}/{ident}.webp"
                expected.add(rel)
                info = inspect_webp(safe_path(ctx.derived, rel))
                if info != entry["outputs"].get(size) or max(info["width"], info["height"]) > limit:
                    report.add(
                        "error", "DERIVED_MISMATCH", "파생 파일 손상/크기/해시 불일치. rebuild 필요", path=rel
                    )
        except Exception:
            report.add("error", "DERIVED_INVALID", "파생 이미지 또는 manifest 검증 실패", record_id=ident)
    if not active_only:
        for path in ctx.derived.rglob("*.webp"):
            if path.relative_to(ctx.derived).as_posix() not in expected:
                report.add(
                    "warning", "UNUSED_DERIVED", "미사용 파생 파일. 자동 삭제하지 않습니다.", path=str(path)
                )
    return manifest


def main():
    p = parser("단계별 검증")
    p.add_argument("--stage", choices=["source", "derived", "public", "remote"], default="source")
    args = p.parse_args()
    ctx = Context(args.root)
    report = Report(ctx, args.stage)
    try:
        if args.stage == "source":
            validate_source(ctx, report)
        elif args.stage == "derived":
            validate_derived(ctx, report)
        elif args.stage == "public":
            from .build_json import validate_public

            validate_public(ctx, report)
        else:
            from .verify_storage import verify

            verify(ctx, report)
    except Exception as exc:
        report.add("error", "VALIDATION_FAILED", f"검증 중 오류 ({type(exc).__name__})")
    report.finish()
    raise SystemExit(0 if report.ok else 1)


if __name__ == "__main__":
    main()
