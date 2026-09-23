"""Preview or apply filename registration and the simplified workbook migration."""
import os
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange
from PIL import Image

from .common import atomic_json, cli, read_json, safe_path, sha
from .validate import validate_source

REMOVED = {
    "04_Morphology_Tree": {"status"},
    "05_Taxon_Morphology": {"presence_status", "documentation_status"},
    "06_Images": {"status"},
    "07_Image_Morphology": {"role", "representative", "display_order"},
}
EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".bmp"}


def rows(ws):
    headers = [c.value for c in ws[1]]
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise ValueError(f"Duplicate headers: {ws.title}")
    return [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in values)]


def append(ws, values):
    ws.append([values.get(c.value) for c in ws[1]])


def delete_column(ws, column):
    """Keep validation targets aligned with retained cells after deletion."""
    for validation in list(ws.data_validations.dataValidation):
        ranges = []
        for area in validation.sqref.ranges:
            start, end = area.min_col, area.max_col
            if start == end == column:
                continue
            start -= int(start > column)
            end -= int(end >= column)
            ranges.append(str(CellRange(min_col=start, max_col=end,
                                        min_row=area.min_row, max_row=area.max_row)))
        if ranges:
            validation.sqref = " ".join(ranges)
        else:
            ws.data_validations.dataValidation.remove(validation)
    ws.delete_cols(column)


def register(ctx, report, apply=False, migrate=False):
    original_hash = sha(ctx.master)
    wb = load_workbook(ctx.master)
    temporary = None
    try:
        changes = 0
        for sheet, columns in REMOVED.items():
            ws = wb[sheet]
            obsolete = [(c.column, c.value) for c in ws[1] if c.value in columns]
            if obsolete and not migrate:
                raise ValueError("Legacy columns remain. Preview with --migrate, then --migrate --apply.")
            for column, name in reversed(obsolete):
                report.add("info", "REMOVE_COLUMN", "Remove obsolete column", sheet=sheet, column=name)
                delete_column(ws, column)
                changes += 1
        taxa = {r["taxon_id"] for r in rows(wb["01_Taxa"])}
        nodes = {r["morphology_id"] for r in rows(wb["04_Morphology_Tree"])}
        images = rows(wb["06_Images"])
        by_path = {}
        by_id = {}
        for row in images:
            path = row["original_filename"]
            if path in by_path or row["image_id"] in by_id:
                raise ValueError("Duplicate original_filename or image_id; reconcile before importing.")
            by_path[path] = row
            by_id[row["image_id"]] = row
            if not safe_path(ctx.originals, path).is_file():
                raise ValueError(f"Registered file missing (possibly renamed): {path}")
        registry_path = ctx.root / ".state/image-id-registry.json"
        registry = read_json(registry_path, {"high_water": 0, "files": {}})
        # Include previous derivative history so deleted IDs are not immediately reused.
        used = set(by_id) | set(read_json(ctx.state, {})) | set(registry["files"])
        high = max([registry["high_water"]] + [int(i[2:]) for i in used if re.fullmatch(r"IM\d+", i)])
        links = {(r["image_id"], r["morphology_id"]) for r in rows(wb["07_Image_Morphology"])}
        states = {(r["taxon_id"], r["morphology_id"]) for r in rows(wb["05_Taxon_Morphology"])}
        known_hashes = {sha(safe_path(ctx.originals, p)): p for p in by_path}
        known_hashes.update({v["sha256"]: v["path"] for v in registry["files"].values()})
        folded = set()
        for path in sorted(ctx.originals.rglob("*")):
            if not path.is_file() or path.name.startswith("."):
                continue
            rel = path.relative_to(ctx.originals).as_posix()
            safe_path(ctx.originals, rel)
            if rel.casefold() in folded:
                raise ValueError(f"Case-colliding original paths: {rel}")
            folded.add(rel.casefold())
            if rel in by_path:
                # Legacy registered names remain valid. Detect conflicts in convention-based names.
                parts = path.stem.split("_", 2)
                if len(parts) == 3 and parts[0] in taxa and parts[1] in nodes:
                    row = by_path[rel]
                    if row["taxon_id"] != parts[0] or (row["image_id"], parts[1]) not in links:
                        raise ValueError(f"Filename conflicts with existing registration: {rel}")
                continue
            if path.suffix.lower() not in EXTENSIONS:
                report.add("warning", "UNSUPPORTED_FILE", "Not imported", path=rel)
                continue
            parts = path.stem.split("_", 2)
            if len(parts) != 3 or not parts[2].strip() or parts[0] not in taxa or parts[1] not in nodes:
                raise ValueError(f"Invalid filename or unknown taxon/morphology: {rel}")
            with Image.open(path) as image:
                image.verify()
            fingerprint = sha(path)
            if fingerprint in known_hashes:
                raise ValueError(f"Possible duplicate/renamed photo: {rel}; registered as {known_hashes[fingerprint]}")
            high += 1
            ident = f"IM{high:06d}"
            row = {"image_id": ident, "taxon_id": parts[0], "original_filename": rel}
            append(wb["06_Images"], row)
            append(wb["07_Image_Morphology"], {"image_id": ident, "morphology_id": parts[1]})
            by_id[ident] = row
            by_path[rel] = row
            links.add((ident, parts[1]))
            known_hashes[fingerprint] = rel
            changes += 2
            report.add("info", "REGISTER_IMAGE", f"{ident}: {parts[0]} / {parts[1]}", path=rel)
        for ident, node in sorted(links):
            if ident not in by_id or node not in nodes:
                raise ValueError("Existing image link references an unknown image or morphology.")
            pair = (by_id[ident]["taxon_id"], node)
            if pair not in states:
                append(wb["05_Taxon_Morphology"], {"taxon_id": pair[0], "morphology_id": node})
                states.add(pair)
                changes += 1
                report.add("info", "REGISTER_MORPHOLOGY", f"{pair[0]} / {node}")
        # Validate a staged workbook, never the live workbook with partial additions.
        fd, filename = tempfile.mkstemp(suffix=".xlsx", dir=ctx.master.parent)
        os.close(fd)
        temporary = Path(filename)
        wb.save(temporary)
        master = ctx.master
        ctx.master = temporary
        try:
            validate_source(ctx, report)
        finally:
            ctx.master = master
        if not report.ok:
            return
        report.add("info", "REGISTRATION_PLAN", f"Changes: {changes}; apply={apply}")
        if not apply:
            return
        if sha(ctx.master) != original_hash:
            raise ValueError("Workbook changed during import; close Excel and retry.")
        if changes:
            backup = ctx.root / "backups" / f"{ctx.master.stem}-{report.run_id}.xlsx"
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ctx.master, backup)
            os.replace(temporary, ctx.master)
            report.add("info", "WORKBOOK_BACKUP", "Original workbook backup", path=str(backup))
        for ident, row in by_id.items():
            registry["files"][ident] = {"path": row["original_filename"],
                                         "sha256": sha(safe_path(ctx.originals, row["original_filename"]))}
        registry["high_water"] = high
        atomic_json(registry_path, registry)
        report.counts["generated"] += changes
    except (ValueError, KeyError, OSError) as error:
        report.add("error", "REGISTRATION_FAILED", str(error))
    finally:
        wb.close()
        if temporary:
            temporary.unlink(missing_ok=True)


def configure(parser):
    parser.add_argument("--apply", action="store_true", help="Save after validation; default is preview only")
    parser.add_argument("--migrate", action="store_true", help="Remove the obsolete columns with backup")


if __name__ == "__main__":
    cli("register", configure, lambda c, r, a: register(c, r, a.apply, a.migrate))
